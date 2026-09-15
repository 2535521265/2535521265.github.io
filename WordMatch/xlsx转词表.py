#!/usr/bin/env python3
"""把 xlsx 单词表转成网页可用的 JS 词表文件。
用法: python xlsx转词表.py 初中.xlsx
输出: 初中.js（与网页同目录即可被读取）
Excel 格式: 第一列单词，第二列答案，首行表头自动跳过。
"""
import sys, json, os
try:
    from openpyxl import load_workbook
except ImportError:
    print('需要先安装 openpyxl: pip install openpyxl')
    sys.exit(1)

def convert(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    pairs = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:  # 跳过表头
            continue
        if not row or len(row) < 2:
            continue
        w = row[0]
        a = row[1]
        if w is None or a is None:
            continue
        w = str(w).strip()
        a = str(a).strip()
        if w and a:
            pairs.append([w, a])
    wb.close()
    name = os.path.splitext(os.path.basename(xlsx_path))[0]
    out_path = os.path.join(os.path.dirname(xlsx_path), name + '.js')
    data = json.dumps(pairs, ensure_ascii=False)
    content = f'window.WORDLISTS=window.WORDLISTS||{{}};window.WORDLISTS[{json.dumps(name,ensure_ascii=False)}]={data};'
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f'已生成 {out_path}（{len(pairs)} 对单词）')

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('用法: python xlsx转词表.py 词表1.xlsx [词表2.xlsx ...]')
        sys.exit(1)
    for p in sys.argv[1:]:
        convert(p)
#（注：内容由AI生成）
